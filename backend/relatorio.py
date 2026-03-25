"""
relatorio.py — Núcleo de geração do relatório de logradouros.
Independente de framework: recebe os arquivos já abertos e retorna os bytes do PDF e Excel.
"""

import io
import os
import hashlib
import zipfile
import glob
import tempfile
import warnings
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.backends.backend_pdf import PdfPages
from PIL import Image as PILImage
import geopandas as gpd
import rasterio
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

warnings.filterwarnings("ignore")

# ── Paleta ──────────────────────────────────────────────────────────────────
CORES = {
    "LEITO NATURAL":           "#E53935",
    "NÃO PAVIMENTADA":         "#E53935",
    "NAO PAVIMENTADA":         "#E53935",
    "SEM PAVIMENTAÇÃO":        "#E53935",
    "PARALELEPIPEDO":          "#43A047",
    "PARALELEPÍPEDO":          "#43A047",
    "INTERTRAVADO":            "#1E88E5",
    "PAVIMENTO INTERTRAVADO":  "#1E88E5",
    "ASFALTO":                 "#0D47A1",
    "REVESTIMENTO ASFALTICO":  "#0D47A1",
    "REVESTIMENTO ASFÁLTICO":  "#0D47A1",
    "PAVIMENTADA":             "#0D47A1",
    "NÃO INFORMADO":           "#9E9E9E",
    "NAO INFORMADO":           "#9E9E9E",
}
CORES_L = ["#E53935", "#43A047", "#1E88E5", "#0D47A1", "#9E9E9E", "#AB47BC", "#FB8C00"]
AMARELO = "#FFD700"
PRETO   = "#2C2C2C"


def get_cor(tipo: str) -> str:
    tu = str(tipo).upper().strip()
    for k, v in CORES.items():
        if k.upper() in tu or tu in k.upper():
            return v
    return CORES_L[int(hashlib.md5(tu.encode()).hexdigest(), 16) % len(CORES_L)]


# ── Helpers ──────────────────────────────────────────────────────────────────
def to_float(v):
    if pd.isna(v):
        return 0.0
    s = str(v).strip().replace(",", ".")
    if s.count(".") > 1:
        p = s.split(".")
        s = "".join(p[:-1]) + "." + p[-1]
    try:
        return float(s)
    except Exception:
        return 0.0


def rodape(fig, nome, data):
    fig.text(0.5, 0.005,
             f"Prefeitura Municipal de {nome} — Relatório de Logradouros — {data}",
             ha="center", fontsize=7, color="#999", style="italic")


def carregar_tabela(file_bytes: bytes, filename: str) -> pd.DataFrame:
    ext = os.path.splitext(filename)[1].lower()
    buf = io.BytesIO(file_bytes)

    if ext in (".xlsx", ".xls"):
        return pd.read_excel(buf, dtype=str)

    if ext == ".dbf":
        from dbfread import DBF
        with tempfile.NamedTemporaryFile(suffix=".dbf", delete=False) as tmp:
            tmp.write(file_bytes)
            tmp_path = tmp.name
        table = DBF(tmp_path, encoding="latin-1", char_decode_errors="ignore")
        df = pd.DataFrame(iter(table))
        os.unlink(tmp_path)
        return df

    # CSV
    for enc in ("utf-8", "latin-1", "cp1252"):
        for sep in (";", ",", "\t"):
            try:
                df = pd.read_csv(io.BytesIO(file_bytes), encoding=enc, sep=sep, dtype=str)
                if len(df.columns) > 1:
                    return df
            except Exception:
                pass
    return pd.read_csv(buf, encoding="latin-1", dtype=str)


def carregar_shapefile(zip_bytes: bytes) -> gpd.GeoDataFrame:
    with tempfile.TemporaryDirectory() as tmpdir:
        zip_buf = io.BytesIO(zip_bytes)
        with zipfile.ZipFile(zip_buf) as z:
            z.extractall(tmpdir)
        shp_files = glob.glob(os.path.join(tmpdir, "**", "*.shp"), recursive=True)
        if not shp_files:
            raise ValueError("Nenhum arquivo .shp encontrado no zip.")
        return gpd.read_file(shp_files[0])


def carregar_tif(tif_bytes: bytes):
    """Retorna (rgb_normalizado, extent, crs)"""
    with tempfile.NamedTemporaryFile(suffix=".tif", delete=False) as tmp:
        tmp.write(tif_bytes)
        tmp_path = tmp.name
    try:
        with rasterio.open(tmp_path) as src:
            rgb = src.read([1, 2, 3]).astype(float)
            bounds = src.bounds
            crs = src.crs
        for b in range(3):
            mn, mx = rgb[b].min(), rgb[b].max()
            if mx > mn:
                rgb[b] = (rgb[b] - mn) / (mx - mn)
        rgb_img = np.transpose(rgb, (1, 2, 0))
        extent = [bounds.left, bounds.right, bounds.bottom, bounds.top]
        return rgb_img, extent, crs
    finally:
        os.unlink(tmp_path)


# ── Gráficos ─────────────────────────────────────────────────────────────────
def grafico_pavimentacao(df_t, df1, nome) -> str:
    pav_km  = df_t.groupby("TIPO")["COMP_TRECH"].sum() / 1000
    pav_via = df1.groupby("TIPO")["RUA"].nunique()
    pav     = pd.DataFrame({"Vias": pav_via, "KM": pav_km}).fillna(0).sort_values("KM", ascending=False)
    cores   = [get_cor(t) for t in pav.index]

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7))
    fig.patch.set_facecolor("white")
    fig.suptitle(f"DISTRIBUIÇÃO POR TIPO DE PAVIMENTAÇÃO\n{nome}", fontsize=15, fontweight="bold", color=PRETO, y=1.01)

    wedges, _, autos = ax1.pie(pav["KM"], autopct="%1.1f%%", colors=cores, startangle=90,
        wedgeprops=dict(width=0.55, edgecolor="white", linewidth=2), pctdistance=0.75)
    for a in autos:
        a.set_fontsize(10); a.set_fontweight("bold")
    ax1.legend(wedges, [f"{t} ({pav.loc[t,'KM']:.1f} km)" for t in pav.index],
        loc="lower center", bbox_to_anchor=(0.5, -0.14), ncol=2, fontsize=9)
    ax1.set_title("Distribuição (km)", fontsize=11)

    bars = ax2.barh(pav.index, pav["KM"], color=cores, edgecolor="white")
    for bar, (idx, row) in zip(bars, pav.iterrows()):
        ax2.text(bar.get_width() + 0.3, bar.get_y() + bar.get_height() / 2,
                 f"{row['KM']:.2f} km  ({row['Vias']:.0f} vias)", va="center", fontsize=9)
    ax2.set_xlabel("Comprimento (km)"); ax2.set_facecolor("#FAFAFA")
    ax2.spines[["top", "right"]].set_visible(False); ax2.invert_yaxis()

    plt.tight_layout()
    path = _tmpimg("g1_pav")
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def grafico_status(df_t, nome) -> str:
    st = df_t.groupby("STATUS")["COMP_TRECH"].sum() / 1000
    cores = []
    for s in st.index:
        su = s.upper()
        if "NÃO" in su or "NAO" in su: cores.append("#E53935")
        elif "ASFALTO" in su or "ASFÁLT" in su: cores.append("#0D47A1")
        else: cores.append("#43A047")

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(14, 6))
    fig.patch.set_facecolor("white")
    fig.suptitle(f"DISTRIBUIÇÃO POR STATUS\n{nome}", fontsize=14, fontweight="bold", color=PRETO)
    ax1.pie(st, autopct="%1.1f%%", colors=cores, startangle=90,
        wedgeprops=dict(width=0.55, edgecolor="white", linewidth=2))
    ax1.legend([f"{s} ({v:.1f} km)" for s, v in st.items()],
        loc="lower center", bbox_to_anchor=(0.5, -0.12), fontsize=9)
    bars = ax2.bar(st.index, st.values, color=cores, edgecolor="white")
    for bar in bars:
        ax2.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.3,
                 f"{bar.get_height():.1f} km", ha="center", fontsize=9, fontweight="bold")
    ax2.set_facecolor("#FAFAFA"); ax2.spines[["top", "right"]].set_visible(False)
    plt.xticks(rotation=15, ha="right"); plt.tight_layout()
    path = _tmpimg("g2_st")
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def grafico_bairros(df_t, df1, nome) -> str:
    bk = df_t.groupby("BAIRRO")["COMP_TRECH"].sum() / 1000
    bv = df1.groupby("BAIRRO")["RUA"].nunique()
    b  = pd.DataFrame({"KM": bk, "Vias": bv}).fillna(0).sort_values("KM", ascending=False).head(15)

    fig, ax = plt.subplots(figsize=(14, 8)); fig.patch.set_facecolor("white")
    bars = ax.barh(b.index, b["KM"], color=AMARELO, edgecolor=PRETO, linewidth=0.5)
    for bar, (idx, row) in zip(bars, b.iterrows()):
        ax.text(bar.get_width() + 0.1, bar.get_y() + bar.get_height() / 2,
                f"{row['KM']:.2f} km  ({row['Vias']:.0f} vias)", va="center", fontsize=8.5)
    ax.set_title(f"TOP 15 BAIRROS POR COMPRIMENTO\n{nome}", fontsize=13, fontweight="bold")
    ax.set_facecolor("#FAFAFA"); ax.spines[["top", "right"]].set_visible(False); ax.invert_yaxis()
    plt.tight_layout()
    path = _tmpimg("g3_bairros")
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def grafico_setores(df_t, df1, nome) -> str:
    sk = df_t.groupby("SETOR")["COMP_TRECH"].sum() / 1000
    sv = df1.groupby("SETOR")["RUA"].nunique()
    s  = pd.DataFrame({"KM": sk, "Vias": sv}).fillna(0).sort_values("KM", ascending=False)
    cores = plt.cm.YlOrRd(np.linspace(0.3, 0.9, len(s)))

    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(16, 7)); fig.patch.set_facecolor("white")
    fig.suptitle(f"DISTRIBUIÇÃO POR SETOR\n{nome}", fontsize=14, fontweight="bold", color=PRETO)
    for ax, col, ylabel in [(ax1, "KM", "Comprimento (km)"), (ax2, "Vias", "Número de Vias")]:
        bars = ax.bar(s.index, s[col], color=cores, edgecolor="white", linewidth=0.5)
        for bar in bars:
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.1,
                    f"{bar.get_height():.1f}" if col == "KM" else str(int(bar.get_height())),
                    ha="center", fontsize=7, fontweight="bold")
        ax.set_ylabel(ylabel); ax.set_facecolor("#FAFAFA")
        ax.spines[["top", "right"]].set_visible(False)
        plt.sca(ax); plt.xticks(rotation=45, ha="right", fontsize=8)
    plt.tight_layout()
    path = _tmpimg("g4_setores")
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def grafico_heatmap(df_t, nome) -> str:
    pivot = df_t.groupby(["SETOR", "TIPO"])["COMP_TRECH"].sum().unstack(fill_value=0) / 1000
    fig, ax = plt.subplots(figsize=(14, 8)); fig.patch.set_facecolor("white")
    im = ax.imshow(pivot.values, cmap="YlOrRd", aspect="auto")
    ax.set_xticks(range(len(pivot.columns))); ax.set_xticklabels(pivot.columns, rotation=30, ha="right", fontsize=9)
    ax.set_yticks(range(len(pivot.index)));   ax.set_yticklabels([f"Setor {s}" for s in pivot.index], fontsize=8)
    for i in range(len(pivot.index)):
        for j in range(len(pivot.columns)):
            v = pivot.values[i, j]
            ax.text(j, i, f"{v:.1f}" if v > 0 else "", ha="center", va="center", fontsize=7,
                    color="white" if v > pivot.values.max() * 0.6 else "#333")
    plt.colorbar(im, ax=ax, label="km", shrink=0.8)
    ax.set_title(f"MAPA DE CALOR: SETOR × TIPO DE PAVIMENTAÇÃO (km)\n{nome}", fontsize=12, fontweight="bold")
    plt.tight_layout()
    path = _tmpimg("g5_heat")
    plt.savefig(path, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def grafico_mapa(gdf: gpd.GeoDataFrame, nome: str, data: str,
                 tif_bytes: bytes | None = None) -> str:
    gdf_plot = gdf.copy()

    if tif_bytes:
        rgb_img, extent, crs_tif = carregar_tif(tif_bytes)
        if gdf_plot.crs != crs_tif:
            gdf_plot = gdf_plot.to_crs(crs_tif)
    else:
        import contextily as cx
        if gdf_plot.crs is None:
            gdf_plot = gdf_plot.set_crs(epsg=4674)
        gdf_plot = gdf_plot.to_crs(epsg=3857)

    fig, ax = plt.subplots(figsize=(14, 12)); fig.patch.set_facecolor("white"); ax.set_facecolor("#CCCCCC")

    if tif_bytes:
        ax.imshow(rgb_img, extent=extent, origin="upper", aspect="auto", zorder=0)
    else:
        basemap_ok = False
        for source in [cx.providers.Esri.WorldImagery, cx.providers.OpenStreetMap.Mapnik]:
            try:
                cx.add_basemap(ax, source=source, zoom="auto", zorder=0)
                basemap_ok = True; break
            except Exception:
                pass
        if not basemap_ok:
            ax.set_facecolor("#B0BEC5")

    tipos = sorted(gdf_plot["TIPO"].dropna().unique())
    patches = []
    for tipo in tipos:
        cor = get_cor(tipo)
        gdf_plot[gdf_plot["TIPO"] == tipo].plot(ax=ax, color=cor, linewidth=1.5, zorder=2, alpha=0.9)
        patches.append(mpatches.Patch(color=cor, label=tipo))

    leg = ax.legend(handles=patches, title="TIPO DE PAVIMENTAÇÃO", loc="lower left",
        fontsize=10, title_fontsize=11, framealpha=0.95, facecolor="white", edgecolor="#AAA")
    leg.get_title().set_fontweight("bold")
    ax.set_title(f"MAPA DE PAVIMENTAÇÃO\n{nome.upper()} — {data}",
        fontsize=16, fontweight="bold", color=PRETO, pad=15)
    ax.tick_params(labelsize=7)
    fig.text(0.5, 0.01, f"Prefeitura Municipal de {nome} — Relatório de Logradouros — {data}",
        ha="center", fontsize=8, color="#9E9E9E", style="italic")
    plt.tight_layout()
    path = _tmpimg("g6_mapa")
    plt.savefig(path, dpi=180, bbox_inches="tight", facecolor="white")
    plt.close()
    return path


def _tmpimg(prefix: str) -> str:
    d = tempfile.mkdtemp()
    return os.path.join(d, f"{prefix}.png")


# ── Excel ────────────────────────────────────────────────────────────────────
def gerar_excel(df_t: pd.DataFrame, df1: pd.DataFrame, nome: str, data: str) -> bytes:
    wb = openpyxl.Workbook()
    FILL_AM = PatternFill("solid", fgColor="FFD700")
    FILL_PK = PatternFill("solid", fgColor="2C2C2C")
    FILL_C  = PatternFill("solid", fgColor="F5F5F5")
    FILL_E  = PatternFill("solid", fgColor="FAFAFA")
    AL_C    = Alignment(horizontal="center", vertical="center", wrap_text=True)
    AL_L    = Alignment(horizontal="left",   vertical="center")
    BORDA   = Border(left=Side(style="thin", color="DDDDDD"), right=Side(style="thin", color="DDDDDD"),
                     top=Side(style="thin", color="DDDDDD"), bottom=Side(style="thin", color="DDDDDD"))

    def hdr(ws, row, n, am=True):
        for c in range(1, n + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = FILL_AM if am else FILL_PK
            cell.font = Font(bold=True, color="2C2C2C" if am else "FFFFFF", size=10)
            cell.alignment = AL_C; cell.border = BORDA

    def linha(ws, row, n, par):
        for c in range(1, n + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = FILL_C if par else FILL_E
            cell.font = Font(color="2C2C2C", size=9)
            cell.alignment = AL_C; cell.border = BORDA

    def autowidth(ws):
        for col in ws.columns:
            l = get_column_letter(col[0].column)
            mx = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[l].width = min(45, max(10, mx + 3))

    total_km = df_t["COMP_TRECH"].sum() / 1000

    # --- Resumo ---
    ws = wb.active; ws.title = "Resumo Geral"; ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:B1"); ws["A1"] = f"RELATÓRIO DE LOGRADOUROS — {nome.upper()}"
    ws["A1"].font = Font(bold=True, size=14, color="2C2C2C"); ws["A1"].fill = FILL_AM; ws["A1"].alignment = AL_C
    ws.merge_cells("A2:B2"); ws["A2"] = f"Gerado em: {data}"
    ws["A2"].font = Font(italic=True, size=9, color="666666"); ws["A2"].alignment = AL_C
    resumo = [("INDICADOR","VALOR"),("Total de Vias",f"{df1['RUA'].nunique():,}"),
              ("Total de Trechos",f"{len(df_t):,}"),("Comprimento Total (km)",f"{total_km:,.2f} km"),
              ("Setores",f"{df1['SETOR'].nunique():,}"),("Bairros",f"{df1['BAIRRO'].nunique():,}")]
    for i,(k,v) in enumerate(resumo,3):
        ws.cell(row=i,column=1).value=k; ws.cell(row=i,column=2).value=v
        if k=="INDICADOR": hdr(ws,i,2,am=False)
        else: linha(ws,i,2,i%2==0)
    autowidth(ws)

    def aba_simples(titulo, grupos):
        ws2 = wb.create_sheet(titulo); ws2.sheet_view.showGridLines = False
        ws2.merge_cells(f"A1:{get_column_letter(len(grupos[0]))}1")
        ws2.cell(1,1).value=f"{titulo.upper()} — {nome.upper()}"
        ws2.cell(1,1).font=Font(bold=True,size=12,color="2C2C2C"); ws2.cell(1,1).fill=FILL_AM; ws2.cell(1,1).alignment=AL_C
        headers=grupos[0]; hdr(ws2,2,len(headers),am=False)
        for j,h in enumerate(headers,1): ws2.cell(2,j).value=h
        for i,row in enumerate(grupos[1:],3):
            for j,v in enumerate(row,1): ws2.cell(i,j).value=v
            linha(ws2,i,len(row),i%2==0)
        autowidth(ws2)

    # Pavimentação
    pav_km=df_t.groupby("TIPO")["COMP_TRECH"].sum()/1000
    pav_v=df1.groupby("TIPO")["RUA"].nunique()
    pav_t=df_t.groupby("TIPO").size()
    pav=pd.DataFrame({"Vias":pav_v,"Trechos":pav_t,"KM":pav_km}).fillna(0).sort_values("KM",ascending=False)
    pav["Pct"]=(pav["KM"]/pav["KM"].sum()*100).round(2)
    rows=[["TIPO","VIAS","TRECHOS","KM","% TOTAL"]]
    for t,r in pav.iterrows(): rows.append([t,int(r.Vias),int(r.Trechos),round(r.KM,3),f"{r.Pct:.2f}%"])
    rows.append(["TOTAL",int(pav.Vias.sum()),int(pav.Trechos.sum()),round(pav.KM.sum(),3),"100%"])
    aba_simples("Pavimentação",rows)

    # Status
    sk=df_t.groupby("STATUS")["COMP_TRECH"].sum()/1000
    sv=df1.groupby("STATUS")["RUA"].nunique()
    st2=df_t.groupby("STATUS").size()
    sdf=pd.DataFrame({"Vias":sv,"Trechos":st2,"KM":sk}).fillna(0).sort_values("KM",ascending=False)
    sdf["Pct"]=(sdf["KM"]/sdf["KM"].sum()*100).round(2)
    rows=[["STATUS","VIAS","TRECHOS","KM","% TOTAL"]]
    for s,r in sdf.iterrows(): rows.append([s,int(r.Vias),int(r.Trechos),round(r.KM,3),f"{r.Pct:.2f}%"])
    aba_simples("Status",rows)

    # Setor
    sk2=df_t.groupby("SETOR")["COMP_TRECH"].sum()/1000
    sv2=df1.groupby("SETOR")["RUA"].nunique()
    sb=df1.groupby("SETOR")["BAIRRO"].nunique()
    st3=df_t.groupby("SETOR").size()
    sdf2=pd.DataFrame({"Bairros":sb,"Vias":sv2,"Trechos":st3,"KM":sk2}).fillna(0).sort_values("KM",ascending=False)
    sdf2["Pct"]=(sdf2["KM"]/sdf2["KM"].sum()*100).round(2)
    rows=[["SETOR","BAIRROS","VIAS","TRECHOS","KM","% TOTAL"]]
    for s,r in sdf2.iterrows(): rows.append([s,int(r.Bairros),int(r.Vias),int(r.Trechos),round(r.KM,3),f"{r.Pct:.2f}%"])
    aba_simples("Por Setor",rows)

    # Bairro
    bk=df_t.groupby("BAIRRO")["COMP_TRECH"].sum()/1000
    bv=df1.groupby("BAIRRO")["RUA"].nunique()
    bt=df_t.groupby("BAIRRO").size()
    bdf=pd.DataFrame({"Vias":bv,"Trechos":bt,"KM":bk}).fillna(0).sort_values("KM",ascending=False)
    bdf["Pct"]=(bdf["KM"]/bdf["KM"].sum()*100).round(2)
    rows=[["BAIRRO","VIAS","TRECHOS","KM","% TOTAL"]]
    for b,r in bdf.iterrows(): rows.append([b,int(r.Vias),int(r.Trechos),round(r.KM,3),f"{r.Pct:.2f}%"])
    aba_simples("Por Bairro",rows)

    # Dados completos
    ws6=wb.create_sheet("Dados Completos"); ws6.sheet_view.showGridLines=False
    cols=["SETOR","BAIRRO","RUA","TRECHO","ID_TRECHO","COMP_TRECH","STATUS","TIPO"]
    hdrs=["SETOR","BAIRRO","LOGRADOURO","TRECHO","ID TRECHO","COMP. (m)","STATUS","TIPO PAVIM."]
    df_exp=df_t[[c for c in cols if c in df_t.columns]].copy()
    ws6.merge_cells(f"A1:{get_column_letter(len(hdrs))}1")
    ws6.cell(1,1).value=f"DADOS COMPLETOS — {nome.upper()}"
    ws6.cell(1,1).font=Font(bold=True,size=12,color="2C2C2C"); ws6.cell(1,1).fill=FILL_AM; ws6.cell(1,1).alignment=AL_C
    for j,h in enumerate(hdrs[:len(df_exp.columns)],1): ws6.cell(2,j).value=h
    hdr(ws6,2,len(df_exp.columns),am=False)
    for i,(_, row) in enumerate(df_exp.iterrows(),3):
        for j,v in enumerate(row,1): ws6.cell(i,j).value=v
        linha(ws6,i,len(df_exp.columns),i%2==0)
    ws6.freeze_panes="A3"; autowidth(ws6)

    buf=io.BytesIO(); wb.save(buf); return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────
def gerar_pdf(df_t, df1, nome, data, img_paths: dict) -> bytes:
    total_km   = df_t["COMP_TRECH"].sum() / 1000
    total_vias = df1["RUA"].nunique()
    total_tr   = len(df_t)

    buf = io.BytesIO()
    with PdfPages(buf) as pdf:

        # Capa
        fig = plt.figure(figsize=(8.5, 11)); fig.patch.set_facecolor("white")
        ax  = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
        ax.add_patch(plt.Rectangle((0, 0.88), 1, 0.12, color=AMARELO, transform=ax.transAxes))
        ax.add_patch(plt.Rectangle((0, 0),    1, 0.06, color=PRETO,   transform=ax.transAxes))
        fig.text(0.5, 0.91, "EIXO ENGENHARIA E PROJETOS", ha="center", fontsize=11, color=PRETO, fontweight="bold")
        fig.text(0.5, 0.65, "RELATÓRIO DE\nLOGRADOUROS", ha="center", fontsize=36, color=PRETO, fontweight="bold", linespacing=1.2)
        fig.text(0.5, 0.56, "Análise Quantitativa da Malha Viária Municipal", ha="center", fontsize=14, color="#555")
        fig.add_artist(plt.Line2D([0.15, 0.85], [0.53, 0.53], transform=fig.transFigure, color=AMARELO, linewidth=3))
        fig.text(0.5, 0.46, nome.upper(), ha="center", fontsize=28, color=PRETO, fontweight="bold")
        fig.text(0.5, 0.41, "Estado da Paraíba", ha="center", fontsize=14, color="#555")
        for val, label, x in [(f"{total_vias}", "VIAS", 0.20), (f"{total_tr:,}", "TRECHOS", 0.50), (f"{total_km:.1f} km", "EXTENSÃO TOTAL", 0.80)]:
            a2 = fig.add_axes([x - 0.10, 0.26, 0.20, 0.10]); a2.axis("off")
            a2.add_patch(plt.Rectangle((0, 0), 1, 1, color=AMARELO))
            fig.text(x, 0.32, val, ha="center", fontsize=20, color=PRETO, fontweight="bold")
            fig.text(x, 0.27, label, ha="center", fontsize=8, color=PRETO)
        fig.text(0.5, 0.20, f"Data de Emissão: {data}", ha="center", fontsize=11, color="#555")
        fig.text(0.5, 0.03, f"Prefeitura Municipal de {nome}", ha="center", fontsize=9, color="white")
        pdf.savefig(fig, bbox_inches="tight", facecolor="white"); plt.close()

        # Visão geral
        fig = plt.figure(figsize=(8.5, 11)); fig.patch.set_facecolor("white")
        ax  = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
        ax.add_patch(plt.Rectangle((0, 0.935), 1, 0.065, color=AMARELO, transform=ax.transAxes))
        fig.text(0.5, 0.953, f"VISÃO GERAL ESTATÍSTICA — {nome.upper()}", ha="center", fontsize=13, color=PRETO, fontweight="bold")
        pav_km  = df_t.groupby("TIPO")["COMP_TRECH"].sum() / 1000
        pav_via = df1.groupby("TIPO")["RUA"].nunique()
        pav     = pd.DataFrame({"Vias": pav_via, "KM": pav_km}).fillna(0).sort_values("KM", ascending=False)
        st_km   = df_t.groupby("STATUS")["COMP_TRECH"].sum() / 1000
        dados   = [["Total de Registros", f"{len(df_t):,}"], ["Total de Vias Únicas", f"{total_vias:,}"],
                   ["Total de Trechos", f"{total_tr:,}"], ["Comprimento Total", f"{total_km:,.2f} km"],
                   ["Número de Setores", f"{df1['SETOR'].nunique():,}"], ["Número de Bairros", f"{df1['BAIRRO'].nunique():,}"],
                   ["", ""], ["TIPO DE PAVIMENTAÇÃO", "COMPRIMENTO (km)"]]
        for tipo in pav.index:
            dados.append([tipo, f"{pav.loc[tipo,'KM']:,.3f} km  ({pav.loc[tipo,'Vias']:.0f} vias)"])
        cell_data = [r for r in dados if r[0]]
        ax_tab = fig.add_axes([0.05, 0.55, 0.90, 0.36]); ax_tab.axis("off")
        tbl = ax_tab.table(cellText=cell_data, loc="upper center", cellLoc="left")
        tbl.auto_set_font_size(False); tbl.set_fontsize(9); tbl.scale(1, 1.6)
        for (r, c), cell in tbl.get_celld().items():
            cell.set_edgecolor("#E0E0E0")
            txt = cell.get_text().get_text()
            if txt in ("TIPO DE PAVIMENTAÇÃO", "COMPRIMENTO (km)"):
                cell.set_facecolor(PRETO); cell.get_text().set_color("white"); cell.get_text().set_fontweight("bold")
            elif r % 2 == 0: cell.set_facecolor("#FFF9E0")
            else: cell.set_facecolor("white")
        cores_p = [get_cor(t) for t in pav.index]
        ax_pie = fig.add_axes([0.05, 0.10, 0.42, 0.38])
        wedges, _, autos = ax_pie.pie(pav["KM"], autopct="%1.1f%%", colors=cores_p, startangle=90,
            wedgeprops=dict(width=0.55, edgecolor="white", linewidth=1.5), pctdistance=0.78)
        for a in autos: a.set_fontsize(8); a.set_fontweight("bold")
        ax_pie.set_title("Pavimentação por km", fontsize=10, fontweight="bold", pad=8)
        ax_pie.legend(wedges, pav.index, loc="lower center", bbox_to_anchor=(0.5, -0.18), fontsize=7, ncol=2)
        cores_st = ["#E53935" if "NÃO" in s.upper() or "NAO" in s.upper() else "#0D47A1" if "ASFALTO" in s.upper() else "#43A047" for s in st_km.index]
        ax_st = fig.add_axes([0.53, 0.10, 0.42, 0.38])
        wedges2, _, autos2 = ax_st.pie(st_km, autopct="%1.1f%%", colors=cores_st, startangle=90,
            wedgeprops=dict(width=0.55, edgecolor="white", linewidth=1.5), pctdistance=0.78)
        for a in autos2: a.set_fontsize(8); a.set_fontweight("bold")
        ax_st.set_title("Status das Vias", fontsize=10, fontweight="bold", pad=8)
        ax_st.legend(wedges2, st_km.index, loc="lower center", bbox_to_anchor=(0.5, -0.18), fontsize=7)
        rodape(fig, nome, data)
        pdf.savefig(fig, bbox_inches="tight", facecolor="white"); plt.close()

        # Gráficos
        paginas = [
            ("g1", "Distribuição por Tipo de Pavimentação"),
            ("g2", "Distribuição por Status"),
            ("g3", "Top 15 Bairros por Comprimento"),
            ("g4", "Distribuição por Setor"),
            ("g5", "Mapa de Calor: Setor × Pavimentação"),
            ("g6", "Mapa de Pavimentação"),
        ]
        for key, titulo in paginas:
            path = img_paths.get(key)
            if not path or not os.path.exists(path):
                continue
            fig = plt.figure(figsize=(8.5, 11)); fig.patch.set_facecolor("white")
            ax  = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
            ax.add_patch(plt.Rectangle((0, 0.935), 1, 0.065, color=AMARELO, transform=ax.transAxes))
            fig.text(0.5, 0.953, titulo.upper(), ha="center", fontsize=12, color=PRETO, fontweight="bold")
            img = PILImage.open(path)
            ai  = fig.add_axes([0.03, 0.07, 0.94, 0.86]); ai.imshow(img); ai.axis("off")
            rodape(fig, nome, data)
            pdf.savefig(fig, bbox_inches="tight", facecolor="white"); plt.close()

        # Tabela setor × tipo
        pivot = df_t.groupby(["SETOR", "TIPO"])["COMP_TRECH"].sum().unstack(fill_value=0) / 1000
        pivot["TOTAL"] = pivot.sum(axis=1); pivot = pivot.sort_values("TOTAL", ascending=False)
        ps = pivot.round(2).reset_index(); ps.columns = [str(c) for c in ps.columns]
        fig = plt.figure(figsize=(8.5, 11)); fig.patch.set_facecolor("white")
        ax  = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
        ax.add_patch(plt.Rectangle((0, 0.935), 1, 0.065, color=AMARELO, transform=ax.transAxes))
        fig.text(0.5, 0.953, "QUANTITATIVO: SETOR × TIPO DE PAVIMENTAÇÃO (km)",
                 ha="center", fontsize=11, color=PRETO, fontweight="bold")
        ax_tab = fig.add_axes([0.02, 0.08, 0.96, 0.84]); ax_tab.axis("off")
        tbl = ax_tab.table(cellText=ps.values.tolist(), colLabels=ps.columns.tolist(),
                           loc="upper center", cellLoc="center")
        tbl.auto_set_font_size(False); tbl.set_fontsize(8); tbl.scale(1, 1.5)
        for (r, c), cell in tbl.get_celld().items():
            cell.set_edgecolor("#E0E0E0")
            if r == 0: cell.set_facecolor(PRETO); cell.get_text().set_color("white"); cell.get_text().set_fontweight("bold")
            elif r % 2 == 0: cell.set_facecolor("#FAFAFA")
            else: cell.set_facecolor("white")
            if c == 0 and r > 0: cell.get_text().set_fontweight("bold")
        rodape(fig, nome, data)
        pdf.savefig(fig, bbox_inches="tight", facecolor="white"); plt.close()

    return buf.getvalue()


# ── Ponto de entrada principal ────────────────────────────────────────────────
def gerar_relatorio(
    tabela_bytes: bytes, tabela_filename: str,
    shp_bytes: bytes,
    nome_municipio: str,
    tif_bytes: bytes | None = None,
) -> tuple[bytes, bytes]:
    """
    Retorna (pdf_bytes, excel_bytes).
    """
    data = datetime.now().strftime("%d/%m/%Y")

    # 1. Carrega dados
    df  = carregar_tabela(tabela_bytes, tabela_filename)
    gdf = carregar_shapefile(shp_bytes)

    # 2. Prepara
    for col in ["SETOR", "BAIRRO", "RUA", "STATUS", "TIPO"]:
        if col in df.columns:
            df[col] = df[col].fillna("NÃO INFORMADO")

    if "COMP_TRECH" in df.columns:
        df["COMP_TRECH"] = df["COMP_TRECH"].apply(to_float)

    df_t = df.copy()
    df1  = df_t[df_t["TRECHO"].apply(lambda x: str(x).strip().split(".")[0] == "1")].copy() if "TRECHO" in df_t.columns else df_t.copy()

    # 3. Gráficos
    imgs = {
        "g1": grafico_pavimentacao(df_t, df1, nome_municipio),
        "g2": grafico_status(df_t, nome_municipio),
        "g3": grafico_bairros(df_t, df1, nome_municipio),
        "g4": grafico_setores(df_t, df1, nome_municipio),
        "g5": grafico_heatmap(df_t, nome_municipio),
        "g6": grafico_mapa(gdf, nome_municipio, data, tif_bytes),
    }

    # 4. PDF + Excel
    pdf_bytes   = gerar_pdf(df_t, df1, nome_municipio, data, imgs)
    excel_bytes = gerar_excel(df_t, df1, nome_municipio, data)

    return pdf_bytes, excel_bytes
